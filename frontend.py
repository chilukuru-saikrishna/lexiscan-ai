import io
import difflib
import requests
import streamlit as st
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

API_URL = "http://127.0.0.1:8000/scan"

st.set_page_config(page_title="Spellbook Clone", layout="wide", initial_sidebar_state="expanded")

st.markdown("""
<style>
    .metric-card {
        background-color: #1e2129;
        border-radius: 8px;
        padding: 16px;
        margin-bottom: 16px;
        border: 1px solid #30363d;
    }
    .diff-box {
        font-family: monospace;
        background-color: #0d1117;
        padding: 12px;
        border-radius: 6px;
        white-space: pre-wrap;
        line-height: 1.6;
        border: 1px solid #30363d;
    }
    .del-text {
        background-color: rgba(248, 81, 73, 0.25);
        color: #ff7b72;
        text-decoration: line-through;
        padding: 2px 4px;
        border-radius: 3px;
    }
    .ins-text {
        background-color: rgba(63, 185, 80, 0.25);
        color: #7ee787;
        font-weight: bold;
        padding: 2px 4px;
        border-radius: 3px;
    }
    .badge-critical { background-color: #da3633; color: white; padding: 3px 8px; border-radius: 12px; font-size: 12px; font-weight: bold; }
    .badge-high { background-color: #d29922; color: black; padding: 3px 8px; border-radius: 12px; font-size: 12px; font-weight: bold; }
    .badge-medium { background-color: #1f6feb; color: white; padding: 3px 8px; border-radius: 12px; font-size: 12px; font-weight: bold; }
    .badge-low { background-color: #238636; color: white; padding: 3px 8px; border-radius: 12px; font-size: 12px; font-weight: bold; }
</style>
""", unsafe_allow_html=True)

def generate_inline_diff(original: str, revised: str) -> str:
    orig_words = original.split()
    rev_words = revised.split()
    matcher = difflib.SequenceMatcher(None, orig_words, rev_words)
    output = []
    
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == 'equal':
            output.append(" ".join(orig_words[i1:i2]))
        elif tag == 'delete':
            output.append(f'<span class="del-text">{" ".join(orig_words[i1:i2])}</span>')
        elif tag == 'insert':
            output.append(f'<span class="ins-text">{" ".join(rev_words[j1:j2])}</span>')
        elif tag == 'replace':
            output.append(f'<span class="del-text">{" ".join(orig_words[i1:i2])}</span>')
            output.append(f'<span class="ins-text">{" ".join(rev_words[j1:j2])}</span>')
            
    return " ".join(output)

def set_cell_background(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_pr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tc_pr.append(tc_mar)

def create_summary_docx(report: dict) -> io.BytesIO:
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
    
    # Title
    title = doc.add_paragraph()
    r_title = title.add_run("LEGAL AUDIT REPORT & CONTRACT REDLINE")
    r_title.bold = True
    r_title.font.size = Pt(20)
    r_title.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
    
    sub = doc.add_paragraph()
    r_sub = sub.add_run(f"Document: {report.get('document_title', 'Master Agreement')} | Perspective: {report.get('reviewed_party', 'Customer')}")
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    
    # --- REPORT CARD SECTION ---
    doc.add_heading("1. Executive Report Card", level=1)
    
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Score Box
    score_cell = table.cell(0, 0)
    score_cell.width = Inches(2.2)
    set_cell_background(score_cell, "F8FAFC")
    set_cell_margins(score_cell, top=140, bottom=140, left=140, right=140)
    p_sc = score_cell.paragraphs[0]
    p_sc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sc_lbl = p_sc.add_run("RISK SCORE\n")
    r_sc_lbl.font.size = Pt(9.5)
    r_sc_lbl.bold = True
    r_sc_lbl.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    
    r_score = p_sc.add_run(f"{report.get('overall_risk_score', 0)} / 10\n")
    r_score.font.size = Pt(32)
    r_score.bold = True
    r_score.font.color.rgb = RGBColor(0xDC, 0x26, 0x26) if report.get('overall_risk_score', 0) >= 7 else RGBColor(0x25, 0x63, 0xEB)
    
    r_sc_status = p_sc.add_run("CRITICAL RISK" if report.get('overall_risk_score', 0) >= 7 else "REVIEW REQUIRED")
    r_sc_status.font.size = Pt(10)
    r_sc_status.bold = True
    
    # Executive Summary Box
    summary_cell = table.cell(0, 1)
    summary_cell.width = Inches(4.6)
    set_cell_background(summary_cell, "F1F5F9")
    set_cell_margins(summary_cell, top=140, bottom=140, left=140, right=140)
    p_sum = summary_cell.paragraphs[0]
    r_sum_title = p_sum.add_run("Executive Assessment:\n")
    r_sum_title.bold = True
    r_sum_title.font.size = Pt(11)
    r_sum_text = p_sum.add_run(report.get("summary", "No summary provided."))
    r_sum_text.font.size = Pt(9.5)
    
    # Missing Clauses Checklist
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    h_missing = doc.add_heading("Critical Missing Protections:", level=2)
    h_missing.paragraph_format.space_before = Pt(12)
    for clause in report.get("missing_clauses", []):
        p_cl = doc.add_paragraph(style='List Bullet')
        r_cl = p_cl.add_run(f"[MISSING]  {clause}")
        r_cl.font.size = Pt(9.5)
        r_cl.font.color.rgb = RGBColor(0xB9, 0x1C, 0x1C)
        r_cl.bold = True
    
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # --- REDLINED & UNDERLINED CONTRACT SECTION ---
    doc.add_heading("2. Clause-by-Clause Audit & Redlines", level=1)
    p_exp = doc.add_paragraph()
    r_exp = p_exp.add_run("The following clauses have been flagged with high contractual risk. Problematic original terms are underlined and highlighted, followed by the proposed revision and negotiation rationale.")
    r_exp.font.size = Pt(9.5)
    r_exp.font.italic = True
    
    for idx, issue in enumerate(report.get("issues", [])):
        sec_ref = issue.get("section_reference", f"Clause #{idx+1}")
        risk = issue.get("risk_level", "HIGH").upper()
        
        # Clause Heading
        h_issue = doc.add_heading(f"{sec_ref} — {issue.get('problem_summary', '')}", level=2)
        h_issue.paragraph_format.space_before = Pt(14)
        h_issue.paragraph_format.space_after = Pt(4)
        
        # Severity Badge
        p_badge = doc.add_paragraph()
        r_b1 = p_badge.add_run(f"SEVERITY: {risk}")
        r_b1.bold = True
        r_b1.font.size = Pt(9)
        r_b1.font.color.rgb = RGBColor(0xDC, 0x26, 0x26) if risk == "CRITICAL" else RGBColor(0xD9, 0x77, 0x06)
        r_b2 = p_badge.add_run(f"  |  CATEGORY: {issue.get('category', 'OTHER')}")
        r_b2.font.size = Pt(9)
        r_b2.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        
        # Original Text (Underlined & Highlighted)
        p_orig = doc.add_paragraph()
        p_orig.paragraph_format.left_indent = Inches(0.2)
        p_orig.paragraph_format.space_after = Pt(4)
        r_orig_lbl = p_orig.add_run("Flagged Original Language:\n")
        r_orig_lbl.bold = True
        r_orig_lbl.font.size = Pt(9.5)
        r_orig_lbl.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        
        r_orig_body = p_orig.add_run(f'"{issue.get("original_text", "")}"')
        r_orig_body.underline = True  # UNDERLINED as requested
        r_orig_body.font.size = Pt(10)
        r_orig_body.font.color.rgb = RGBColor(0x99, 0x1B, 0x1B)  # Red text
        
        # Redline Revision Box
        redline_tbl = doc.add_table(rows=1, cols=1)
        redline_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        redline_tbl.autofit = False
        
        cell = redline_tbl.cell(0, 0)
        cell.width = Inches(6.8)
        set_cell_background(cell, "F0FDF4")  # Soft green highlight
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        
        p_rev = cell.paragraphs[0]
        r_rev_lbl = p_rev.add_run("Proposed Replacement Language:\n")
        r_rev_lbl.bold = True
        r_rev_lbl.font.size = Pt(9.5)
        r_rev_lbl.font.color.rgb = RGBColor(0x16, 0x65, 0x34)
        
        r_rev_body = p_rev.add_run(issue.get("suggested_revision", ""))
        r_rev_body.font.size = Pt(10)
        r_rev_body.font.color.rgb = RGBColor(0x15, 0x80, 0x3D)
        
        # Counsel Rationale
        p_rat = doc.add_paragraph()
        p_rat.paragraph_format.left_indent = Inches(0.2)
        p_rat.paragraph_format.space_before = Pt(4)
        p_rat.paragraph_format.space_after = Pt(12)
        r_rat_lbl = p_rat.add_run("Negotiation Rationale / Comment: ")
        r_rat_lbl.bold = True
        r_rat_lbl.font.size = Pt(9.5)
        r_rat_lbl.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        
        r_rat_body = p_rat.add_run(issue.get("rationale", ""))
        r_rat_body.font.size = Pt(9.5)
        r_rat_body.font.italic = True
        r_rat_body.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream

st.title("⚡ Spellbook Review Engine")
st.caption("Autonomous contract risk analysis & redline generator")

with st.sidebar:
    st.header("Scan Settings")
    uploaded_file = st.file_uploader("Upload Agreement", type=["pdf", "docx", "txt"])
    party = st.selectbox("Your Position", ["Customer", "Buyer", "Vendor", "Licensee", "Service Provider"])
    playbook = st.text_area(
        "Playbook Guidelines",
        value="Standard corporate protections. Cap liability to 12 months fees paid. Mutual indemnification required. Retain IP ownership of custom deliverables. Bilateral termination rights.",
        height=120
    )
    scan_button = st.button("🚀 Analyze Contract", type="primary", use_container_width=True)

if scan_button and uploaded_file is not None:
    with st.spinner("Analyzing document and generating redlines..."):
        try:
            files = {"file": (uploaded_file.name, uploaded_file.getvalue(), uploaded_file.type)}
            data = {"party": party, "playbook": playbook}
            
            response = requests.post(API_URL, files=files, data=data)
            
            if response.status_code == 200:
                st.session_state["report"] = response.json()
            else:
                st.error(f"API Error ({response.status_code}): {response.text}")
        except Exception as e:
            st.error(f"Failed to connect to backend server: {str(e)}")

# DISPLAY RESULTS & DOWNLOAD BUTTON
if "report" in st.session_state:
    report = st.session_state["report"]
    
    # --- DOWNLOAD BUTTON DIRECTLY UNDER SIDEBAR / TOP ---
    with st.sidebar:
        st.markdown("---")
        st.subheader("📄 Export Report")
        docx_data = create_summary_docx(report)
        st.download_button(
            label="📥 Download Legal Audit Report (.docx)",
            data=docx_data,
            file_name=f"{report.get('document_title', 'Contract')}_Audit_Report.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            type="secondary",
            use_container_width=True
        )
    
    st.markdown("---")
    
    col1, col2, col3 = st.columns([1, 2, 2])
    with col1:
        st.markdown(f"""
        <div class="metric-card" style="text-align: center;">
            <div style="font-size: 14px; color: #8b949e;">OVERALL RISK SCORE</div>
            <div style="font-size: 48px; font-weight: bold; color: {'#f85149' if report['overall_risk_score'] >= 7 else '#d29922' if report['overall_risk_score'] >= 4 else '#3fb950'};">
                {report['overall_risk_score']} / 10
            </div>
            <div style="font-size: 12px; color: #8b949e;">Party: {report.get('reviewed_party', 'N/A')}</div>
        </div>
        """, unsafe_allow_html=True)
        
    with col2:
        st.markdown(f"""
        <div class="metric-card">
            <strong style="color: #58a6ff;">Executive Summary</strong><br>
            <p style="font-size: 13px; color: #c9d1d9; margin-top: 8px;">{report.get('summary', 'No summary provided.')}</p>
        </div>
        """, unsafe_allow_html=True)
        
    with col3:
        st.markdown("""<div class="metric-card"><strong style="color: #f85149;">Missing Standard Clauses</strong><ul style="font-size: 13px; color: #c9d1d9; margin-top: 8px; padding-left: 20px;">""" + 
                    "".join([f"<li>{clause}</li>" for clause in report.get("missing_clauses", [])]) + 
                    "</ul></div>", unsafe_allow_html=True)

    st.subheader(f"Clause Redlines & Findings ({len(report.get('issues', []))})")
    
    for idx, issue in enumerate(report.get("issues", [])):
        risk = issue.get("risk_level", "MEDIUM").upper()
        badge_class = f"badge-{risk.lower()}"
        
        with st.expander(f"{issue.get('section_reference', f'Issue #{idx+1}')} — {issue.get('problem_summary', '')}", expanded=True):
            st.markdown(f"""
            <div style="margin-bottom: 12px;">
                <span class="{badge_class}">{risk}</span> &nbsp; 
                <span style="color: #8b949e; font-size: 12px;">CATEGORY: <b>{issue.get('category', 'OTHER')}</b></span>
            </div>
            """, unsafe_allow_html=True)
            
            diff_html = generate_inline_diff(issue.get("original_text", ""), issue.get("suggested_revision", ""))
            
            st.markdown("<b>Track Changes / Redline View:</b>", unsafe_allow_html=True)
            st.markdown(f'<div class="diff-box">{diff_html}</div>', unsafe_allow_html=True)
            
            col_a, col_b = st.columns(2)
            with col_a:
                st.markdown("<p style='font-size: 12px; color: #8b949e; margin-top: 10px;'><b>Original Text:</b></p>", unsafe_allow_html=True)
                st.info(issue.get("original_text", ""))
            with col_b:
                st.markdown("<p style='font-size: 12px; color: #8b949e; margin-top: 10px;'><b>Proposed Language:</b></p>", unsafe_allow_html=True)
                st.success(issue.get("suggested_revision", ""))
                
            st.markdown(f"<b>Negotiation Rationale:</b> {issue.get('rationale', '')}")